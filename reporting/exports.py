"""
Downloads for a class: class list, grades and attendance as Excel
workbooks, and finalized reports as one PDF (a page per student).

Callers pass in students already limited to what the requester may see
(accounts.scoping), so this module never decides who sees what.
"""
from collections import defaultdict
from io import BytesIO
from pathlib import Path

import openpyxl
from django.utils import timezone
from fpdf import FPDF
from fpdf.fonts import FontFace
from openpyxl.styles import Alignment, Font, PatternFill

from attendance.models import AttendanceRecord
from gradebook.levels import level_for, levels_key, with_level
from gradebook.models import Grade
from gradebook.weighting import school_weights, subject_percents
from students.presets import DEFAULT_VOCAB, words_for

from .models import StudentReport
from .spreadsheets import append_row

HEADER_FILL = PatternFill("solid", fgColor="1E2F52")
HEADER_FONT = Font(bold=True, color="FAF7F0")
FONT_DIR = Path(__file__).resolve().parent / "fonts"
GENDERS = {"female": "Female", "male": "Male", "other": "Other"}
MODES = {"day": "Day", "boarding": "Boarding"}


def _name(student):
    return f"{student.first_name} {student.last_name}"


def _sheet(wb, title, headers, rows, widths=None):
    ws = wb.create_sheet(title)
    append_row(ws, headers)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        append_row(ws, row)
    ws.freeze_panes = "A2"
    for i, header in enumerate(headers, start=1):
        width = (widths or {}).get(header) or max(10, min(40, len(str(header)) + 4))
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    return ws


def _workbook_bytes(wb):
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def class_list_xlsx(students, words=DEFAULT_VOCAB):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    rows = []
    for s in students:
        parents = [g for g in s.guardians.all() if g.user.is_active]
        rows.append([
            s.external_id, s.last_name, s.first_name, GENDERS.get(s.gender, ""), s.date_of_birth,
            s.house, MODES.get(s.mode_of_learning, ""), s.enrolled_on, "Active" if s.is_active else "Inactive",
            "; ".join(g.name for g in parents), "; ".join(g.user.email for g in parents),
            "; ".join(g.phone for g in parents if g.phone), s.medical_notes,
        ])
    ws = _sheet(wb, "Class list", [
        words["student_id"], "Last name", "First name", "Gender", "Date of birth", "House", "Mode of learning",
        "Admission date", "Status", "Parents", "Parent emails", "Parent phones", "Health notes",
    ], rows, widths={"Parents": 28, "Parent emails": 32, "Parent phones": 22, "Health notes": 40})
    for row in ws.iter_rows(min_row=2):
        for cell in (row[4], row[7]):
            cell.number_format = "yyyy-mm-dd"
    return _workbook_bytes(wb)


def grades_xlsx(students, term, scale="percent", words=DEFAULT_VOCAB):
    grades = Grade.objects.filter(student__in=students, term=term).select_related("subject")
    by_student = defaultdict(dict)
    subjects = set()
    for (student_id, _term), per in subject_percents(grades, school_weights(term.school),
                                                     key=lambda g: g.subject.name).items():
        for subject, value in per.items():
            by_student[student_id][subject] = round(value, 1)
            subjects.add(subject)
    subjects = sorted(subjects)
    rows = []
    for s in students:
        marks = [by_student[s.id].get(subj) for subj in subjects]
        present = [m for m in marks if m is not None]
        average = round(sum(present) / len(present), 1) if present else None
        row = [s.external_id, s.last_name, s.first_name, *marks, average]
        if level_for(0, scale):  # CBC schools also get the level for the average
            row.append(level_for(average, scale))
        rows.append(row)
    headers = [words["student_id"], "Last name", "First name", *subjects, "Average"]
    if level_for(0, scale):
        headers.append("Level")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Grades (%)", headers, rows)
    return _workbook_bytes(wb)


def attendance_xlsx(students, start, end, words=DEFAULT_VOCAB):
    records = AttendanceRecord.objects.filter(student__in=students, date__gte=start, date__lte=end)
    by_student = defaultdict(dict)
    days = set()
    for r in records:
        by_student[r.student_id][r.date] = r.status
        days.add(r.date)
    days = sorted(days)
    letters = {"present": "P", "absent": "A", "late": "L", "excused": "E"}

    summary = []
    for s in students:
        statuses = list(by_student[s.id].values())
        counts = {k: statuses.count(k) for k in letters}
        total = len(statuses)
        rate = round((counts["present"] + counts["late"]) / total * 100, 1) if total else None
        summary.append([s.external_id, s.last_name, s.first_name, total, counts["present"], counts["absent"],
                        counts["late"], counts["excused"], rate])
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Summary", [words["student_id"], "Last name", "First name", "Days recorded", "Present", "Absent",
                           "Late", "Excused", "Attendance % (present or late)"],
           summary, widths={"Attendance % (present or late)": 30})
    grid = [[s.external_id, _name(s), *[letters.get(by_student[s.id].get(d), "") for d in days]] for s in students]
    ws = _sheet(wb, "By day", [words["student_id"], "Student", *[d.strftime("%a %d %b") for d in days]], grid,
                widths={"Student": 26})
    append_row(ws, [])
    append_row(ws, ["P = present, A = absent, L = late, E = excused"])
    return _workbook_bytes(wb)


class _ReportPDF(FPDF):
    def __init__(self, school_name):
        super().__init__(format="A4")
        self.school_name = school_name
        self.add_font("Serif", "", str(FONT_DIR / "DejaVuSerif.ttf"))
        self.add_font("Serif", "B", str(FONT_DIR / "DejaVuSerif-Bold.ttf"))
        self.set_auto_page_break(auto=True, margin=18)
        self.set_margins(18, 16, 18)

    def footer(self):
        self.set_y(-12)
        self.set_font("Serif", "", 8)
        self.set_text_color(110, 110, 110)
        self.cell(0, 6, f"{self.school_name} · Page {self.page_no()}", align="C")


NAVY, GOLD, INK, MUTED = (30, 47, 82), (184, 134, 46), (27, 35, 51), (110, 110, 110)
HEAD_STYLE = FontFace(emphasis="BOLD", fill_color=(242, 237, 225))


def _pct(value):
    return f"{value:.0f}%" if value is not None else "—"


def _results_table(pdf, school, summary, words):
    """The results table for the school's system. Returns the summary line under it."""
    system, rows = summary["system"], summary["subjects"]
    subject = words["subject"]
    if system == "844":
        head = [subject, "Marks", "Grade", "Points", "Remarks"]
        body = [[r["subject"], _pct(r["percent"]), r["kcse_grade"] or "—", r["points"] or "—", r["comment"]] for r in rows]
        widths = (44, 20, 18, 18, 74)
    elif system == "american":
        head = [subject, "Percent", "Grade", "Credits", "Points", "Comment"]
        body = [[r["subject"], _pct(r["percent"]), r["letter"] or "—", f"{r['credits']:g}",
                 r["gpa_points"] if r["gpa_points"] is not None else "—", r["comment"]] for r in rows]
        widths = (40, 20, 16, 17, 16, 65)
    elif system == "ib" and not any(r["criteria"] for r in rows):
        # Diploma Programme students are graded 1 to 7 from their marks, without MYP criteria.
        head = [subject, "Percent", "Grade", "Comment"]
        body = [[f'{r["subject"]} {r["subject_level"]}'.strip(), _pct(r["percent"]), r["ib_grade"] or "—", r["comment"]]
                for r in rows]
        widths = (50, 20, 18, 86)
    elif system == "ib":
        head = [subject, "A", "B", "C", "D", "Grade", "Comment"]
        body = [[f'{r["subject"]} {r["subject_level"]}'.strip(), *[r["criteria"].get(c, "—") for c in "ABCD"],
                 r["ib_grade"] or "—", r["comment"]] for r in rows]
        widths = (42, 11, 11, 11, 11, 16, 72)
    elif system == "british":
        head = [subject, "Percent", "Grade", "Effort", "Target", "Comment"]
        body = [[r["subject"], _pct(r["percent"]), r["level"] or "—", r["effort"] or "—", r["target"] or "—",
                 r["comment"]] for r in rows]
        widths = (40, 18, 16, 16, 16, 68)
    elif system == "cbc":
        head = [subject, "Score", "Level", "Teacher's comment"]
        body = [[r["subject"], _pct(r["percent"]), r["level"] or "—", r["comment"]] for r in rows]
        widths = (50, 18, 18, 88)
    else:
        leveled = bool(level_for(0, summary["scale"]))
        head = [subject, "Percent"] + (["Level"] if leveled else []) + ["Comment"]
        body = [[r["subject"], _pct(r["percent"])] + ([r["level"] or "—"] if leveled else []) + [r["comment"]]
                for r in rows]
        widths = (50, 20, 18, 86) if leveled else (56, 22, 96)
    align = ("LEFT",) + ("CENTER",) * (len(widths) - 2) + ("LEFT",)
    with pdf.table(col_widths=widths, text_align=align, line_height=6, headings_style=HEAD_STYLE) as table:
        table.row([str(h) for h in head])
        for row in body:
            table.row([str(c) if c not in (None, "") else "" for c in row])

    if system == "844" and summary.get("mean_grade"):
        line = (f"Total marks: {summary['total_marks']}   ·   Total points: {summary['total_points']}   ·   "
                f"Mean grade: {summary['mean_grade']} ({summary['mean_points']:g} points)")
        positions = summary.get("positions")
        if positions:
            line += (f"\nPosition in {words['class'].lower()}: {positions['stream']['position']} of "
                     f"{positions['stream']['of']}   ·   Position in {words['year_group'].lower()}: "
                     f"{positions['form']['position']} of {positions['form']['of']}")
        return line
    if system == "american" and summary.get("gpa") is not None:
        extra = f"   ·   {summary['honor_roll']}" if summary["honor_roll"] else ""
        return f"GPA: {summary['gpa']:.2f} (4.0 scale, {summary['credits']:g} credits){extra}"
    if system == "ib" and summary.get("ib_total") is not None:
        return f"Total of grades: {summary['ib_total']}"
    if summary.get("average") is not None:
        return f"Overall average: {with_level(summary['average'], summary['scale'])}"
    return ""


def _ratings(pdf, report, system):
    """CBC competencies and values, or IB approaches to learning, when the teacher has rated them."""
    from gradebook.systems import REPORT_EXTRAS

    for group in REPORT_EXTRAS.get(system, []):
        given = (report.extra or {}).get(group["key"]) or {}
        if not given:
            continue
        names = group.get("rating_names", {})
        pdf.set_font("Serif", "B", 11)
        pdf.set_text_color(*NAVY)
        pdf.cell(0, 7, group["title"], new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Serif", "", 9)
        pdf.set_text_color(*INK)
        items = [item for item in group["items"] if given.get(item)]
        widths = (52, 35, 52, 35) if names else (62, 25, 62, 25)
        with pdf.table(col_widths=widths, text_align=("LEFT", "CENTER", "LEFT", "CENTER"),
                       line_height=5.5, first_row_as_headings=False) as table:
            for i in range(0, len(items), 2):
                pair = items[i:i + 2]
                cells = []
                for item in pair:
                    rating = given[item]
                    cells += [item, f"{rating} ({names[rating]})" if rating in names else rating]
                table.row(cells + ["", ""] * (2 - len(pair)))
        pdf.ln(3)


def _key(system, scale):
    parts = []
    if key := levels_key(scale):
        parts.append(f"Levels: {key}")
    if system == "844":
        parts.append("Points: A = 12, A- = 11, B+ = 10 … D- = 2, E = 1; mean grade from mean points")
    elif system == "ib":
        parts.append("MYP grade from criteria A to D (each out of 8) when all four are assessed")
    elif system == "american":
        parts.append("Grade points: A = 4, B = 3, C = 2, D = 1, F = 0, weighted by credits")
    return parts


def reports_pdf(school, students, term):
    """One page per student with a finalized report for the term, in the school's system. Returns (bytes, count)."""
    from gradebook.systems import term_summary

    reports = {
        r.student_id: r
        for r in StudentReport.objects.filter(student__in=students, term=term, status="finalized")
    }
    attendance = defaultdict(lambda: defaultdict(int))
    if term.start_date and term.end_date:
        for r in AttendanceRecord.objects.filter(student__in=students, date__gte=term.start_date,
                                                 date__lte=term.end_date):
            attendance[r.student_id][r.status] += 1

    pdf = _ReportPDF(school.name)
    count = 0
    for s in students:
        report = reports.get(s.id)
        if report is None:
            continue
        s.school = school  # the caller's copy has the current settings
        count += 1
        pdf.add_page()
        pdf.set_draw_color(*GOLD)
        pdf.set_line_width(1.2)
        pdf.line(18, 14, 192, 14)
        pdf.set_line_width(0.2)
        pdf.set_draw_color(221, 213, 194)
        pdf.set_text_color(*NAVY)
        pdf.set_font("Serif", "B", 11)
        pdf.cell(0, 8, school.name.upper(), new_x="LMARGIN", new_y="NEXT")
        # The school's motto and contact details, when set during setup.
        contact = " · ".join(x for x in (" ".join(school.address.split()), school.phone, school.email) if x)
        if school.motto or contact:
            pdf.set_font("Serif", "", 8)
            pdf.set_text_color(90, 90, 90)
            for line in (school.motto, contact):
                if line:
                    pdf.cell(0, 4.5, line, new_x="LMARGIN", new_y="NEXT")
            pdf.set_text_color(*NAVY)
            pdf.ln(1)
        pdf.set_font("Serif", "B", 20)
        pdf.cell(0, 11, _name(s), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Serif", "", 10.5)
        pdf.set_text_color(*INK)
        klass = s.school_class
        summary = term_summary(s, term)
        words = words_for(school, summary["system"])
        details = [f"{words['term']}: {term.name}"]
        if klass:
            details.append(f"{words['class']}: {klass.year_group.name} · {klass.name}")
        if s.external_id:
            details.append(f"{words['student_id']}: {s.external_id}")
        if s.pathway:
            details.append(f"Pathway: {s.pathway}")
        pdf.cell(0, 7, "   |   ".join(details), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3)

        # A school running two systems prints each student's card in their own section's style.
        if summary["subjects"]:
            pdf.set_font("Serif", "B", 12)
            pdf.set_text_color(*NAVY)
            pdf.cell(0, 8, "Results", new_x="LMARGIN", new_y="NEXT")
            pdf.set_text_color(*INK)
            pdf.set_font("Serif", "", 9)
            line = _results_table(pdf, school, summary, words)
            if line:
                pdf.ln(2)
                pdf.set_font("Serif", "B", 10)
                pdf.multi_cell(0, 6, line, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(3)

        _ratings(pdf, report, summary["system"])

        att = attendance[s.id]
        total = sum(att.values())
        if total:
            pdf.set_font("Serif", "", 10)
            rate = (att["present"] + att["late"]) / total * 100
            pdf.cell(0, 7, f"Attendance this {words['term'].lower()}: {rate:.0f}% ({att['absent']} absent, "
                           f"{att['late']} late, {att['excused']} excused, of {total} days)",
                     new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2)

        for title, text in (("Class teacher's comment", report.report_comment),
                            ("Principal's remarks", report.principal_comment)):
            if not text and title.startswith("Principal"):
                continue
            pdf.set_font("Serif", "B", 11)
            pdf.set_text_color(*NAVY)
            pdf.cell(0, 7, title, new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("Serif", "", 10.5)
            pdf.set_text_color(*INK)
            pdf.multi_cell(0, 6, text or "—", align="L", new_x="LMARGIN", new_y="NEXT")
            pdf.ln(3)

        pdf.ln(2)
        pdf.set_font("Serif", "", 7.5)
        pdf.set_text_color(*MUTED)
        for part in _key(summary["system"], summary["scale"]):
            pdf.multi_cell(0, 4.5, part, new_x="LMARGIN", new_y="NEXT")
        finalized = report.finalized_at.date() if report.finalized_at else timezone.localdate()
        pdf.cell(0, 5, f"Finalized {finalized:%d %B %Y}", new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output()), count
