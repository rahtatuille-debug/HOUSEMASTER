"""
Student rankings: one rule everywhere. Tied students share a position
(1, 2, 2, 4); 8-4-4 ranks by mean points then average mark (as its report
card does), other curricula by average; CBC doesn't rank learners.
"""
import io
from datetime import date

import openpyxl

from accounts.tests import SchoolScopedAPITestCase
from gradebook.models import Grade, Subject, Term
from gradebook.systems import term_summary
from reporting.rankings import positions
from students.models import School, SchoolClass, Student, YearGroup


class PositionsTests(SchoolScopedAPITestCase):
    def test_ties_share_a_position_and_the_next_one_skips(self):
        self.assertEqual(positions({1: (80,), 2: (70,), 3: (70,), 4: (60,)}), {1: 1, 2: 2, 3: 2, 4: 4})

    def test_students_without_a_score_are_not_ranked(self):
        self.assertEqual(positions({1: None, 2: (50,)}), {2: 1})


class RankingFixture(SchoolScopedAPITestCase):
    """A school whose own curriculum is set per test, with two classes in one year group and two terms."""

    system, scale = "british", "igcse9"

    def setUp(self):
        super().setUp()
        School.objects.filter(pk=self.school_a.pk).update(education_system=self.system, grading_scale=self.scale)
        self.school_a.refresh_from_db()
        self.admin = self.authed_client(self.admin_a)
        self.year = YearGroup.objects.create(school=self.school_a, name="Form 3")
        self.east = SchoolClass.objects.create(year_group=self.year, name="East")
        self.west = SchoolClass.objects.create(year_group=self.year, name="West")
        self.t1 = Term.objects.create(school=self.school_a, name="T1", start_date=date(2026, 1, 1))
        self.t2 = Term.objects.create(school=self.school_a, name="T2", start_date=date(2026, 5, 1))
        self.maths = Subject.objects.create(school=self.school_a, name="Maths")
        self.english = Subject.objects.create(school=self.school_a, name="English")
        self.chem = Subject.objects.create(school=self.school_a, name="Chemistry")

    def pupil(self, name, klass, t2, t1=None):
        s = Student.objects.create(school=self.school_a, first_name=name, last_name="K", school_class=klass)
        for term, marks in ((self.t2, t2), (self.t1, t1 or {})):
            for subject, score in marks.items():
                Grade.objects.create(student=s, subject=subject, term=term, score=score)
        return s

    def perf(self, client=None, **params):
        return (client or self.admin).get("/api/analytics/performance/", {"term": self.t2.id, **params}).data

    def rows(self, **params):
        return {r["name"].split()[0]: r for r in self.perf(**params)["students"]}


class AverageRankingTests(RankingFixture):
    def setUp(self):
        super().setUp()
        # T2 averages: Ann 80, Ben 70, Cat 70, Dan 60. T1: Ann 85, Ben 50, Cat 69, Dan 30.
        self.ann = self.pupil("Ann", self.east, {self.maths: 80, self.english: 80}, {self.maths: 85})
        self.ben = self.pupil("Ben", self.east, {self.maths: 60, self.english: 80}, {self.maths: 50})
        self.cat = self.pupil("Cat", self.east, {self.maths: 90, self.english: 50}, {self.maths: 69})
        self.dan = self.pupil("Dan", self.west, {self.maths: 60, self.chem: 60}, {self.maths: 30})

    def test_class_positions_share_ties(self):
        rows = self.rows(scope="class", id=self.east.id)
        self.assertEqual({n: (r["position"], r["of"]) for n, r in rows.items()},
                         {"Ann": (1, 3), "Ben": (2, 3), "Cat": (2, 3)})

    def test_rank_by_subject_only_counts_students_who_take_it(self):
        rows = self.rows(scope="year_group", id=self.year.id)
        self.assertEqual(rows["Cat"]["subject_positions"]["Maths"], 1)
        self.assertEqual(rows["Ann"]["subject_positions"]["Maths"], 2)
        self.assertEqual({rows["Ben"]["subject_positions"]["Maths"], rows["Dan"]["subject_positions"]["Maths"]}, {3})
        self.assertEqual(rows["Dan"]["subject_positions"]["Chemistry"], 1)
        self.assertNotIn("Chemistry", rows["Ann"]["subject_positions"])
        self.assertEqual(rows["Ann"]["subjects"]["English"], 80.0)
        self.assertEqual(rows["Ann"]["subject_of"]["Maths"], 4)

    def test_most_improved(self):
        # Changes since T1: Dan +30, Ben +20, Cat +1, Ann -5.
        rows = self.rows(scope="year_group", id=self.year.id)
        self.assertEqual({n: r["improvement_position"] for n, r in rows.items()},
                         {"Dan": 1, "Ben": 2, "Cat": 3, "Ann": 4})

    def test_whole_school_ranking_for_admins(self):
        data = self.perf(scope="school")
        rows = {r["name"].split()[0]: r for r in data["students"]}
        self.assertEqual(rows["Ann"]["position"], 1)
        self.assertEqual(rows["Dan"]["position"], 4)
        self.assertEqual(rows["Dan"]["class_name"], "West")
        self.assertEqual(rows["Dan"]["year_group_name"], "Form 3")

    def test_teachers_see_their_own_students_ranked_against_the_whole_group(self):
        self.assign(self.user_a, self.west, self.maths)
        teacher = self.authed_client(self.user_a)
        rows = self.perf(teacher, scope="year_group", id=self.year.id)["students"]
        self.assertEqual([r["name"] for r in rows], ["Dan K"])
        self.assertEqual((rows[0]["position"], rows[0]["of"]), (4, 4))
        self.assertEqual(self.perf(teacher, scope="school").get("students"), None)

    def test_export_has_positions_with_ties(self):
        response = self.admin.get("/api/exports/grades/", {"school_class": self.east.id, "term": self.t2.id})
        sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
        header = [c.value for c in sheet[1]]
        self.assertIn("Position", header)
        col = header.index("Position")
        by_name = {row[header.index("First name")]: row[col] for row in sheet.iter_rows(min_row=2, values_only=True)}
        self.assertEqual(by_name, {"Ann": 1, "Ben": 2, "Cat": 2})


class KcseRankingTests(RankingFixture):
    system, scale = "844", "kcse"

    def setUp(self):
        super().setUp()
        # Xen: 80, 80 -> A, A = 12 mean points, average 80.
        # Yus: 100, 79 -> A, A- = 11.5 mean points, average 89.5.
        # By average Yus would be first; by mean points (the KCSE rule) Xen is.
        self.xen = self.pupil("Xen", self.east, {self.maths: 80, self.english: 80})
        self.yus = self.pupil("Yus", self.east, {self.maths: 100, self.english: 79})

    def test_ranked_by_mean_points_like_the_report_card(self):
        rows = self.rows(scope="class", id=self.east.id)
        self.assertEqual((rows["Xen"]["position"], rows["Yus"]["position"]), (1, 2))
        self.assertEqual(rows["Xen"]["mean_grade"], "A")
        self.assertEqual(rows["Yus"]["mean_points"], 11.5)
        for student in (self.xen, self.yus):
            card = term_summary(student, self.t2)["positions"]["stream"]["position"]
            self.assertEqual(rows[student.first_name]["position"], card)

    def test_export_adds_total_marks_points_and_mean_grade(self):
        response = self.admin.get("/api/exports/grades/", {"school_class": self.east.id, "term": self.t2.id})
        sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
        header = [c.value for c in sheet[1]]
        for column in ("Total marks", "Total points", "Mean grade", "Position"):
            self.assertIn(column, header)
        xen = next(r for r in sheet.iter_rows(min_row=2, values_only=True) if r[header.index("First name")] == "Xen")
        self.assertEqual(xen[header.index("Total marks")], 160)
        self.assertEqual(xen[header.index("Total points")], 24)
        self.assertEqual(xen[header.index("Mean grade")], "A")
        self.assertEqual(xen[header.index("Position")], 1)


class CbcIsNotRankedTests(RankingFixture):
    system, scale = "cbc", "cbc"

    def setUp(self):
        super().setUp()
        self.pupil("Amo", self.east, {self.maths: 90}, {self.maths: 40})
        self.pupil("Bea", self.east, {self.maths: 50}, {self.maths: 60})

    def test_no_positions_of_any_kind(self):
        for row in self.perf(scope="class", id=self.east.id)["students"]:
            self.assertIsNone(row["position"])
            self.assertIsNone(row["improvement_position"])
            self.assertEqual(row["subject_positions"], {})

    def test_export_has_no_position_column(self):
        response = self.admin.get("/api/exports/grades/", {"school_class": self.east.id, "term": self.t2.id})
        header = [c.value for c in openpyxl.load_workbook(io.BytesIO(response.content)).active[1]]
        self.assertNotIn("Position", header)


class MixedCurriculaTests(RankingFixture):
    system, scale = "844", "kcse"

    def setUp(self):
        super().setUp()
        # A CBC junior section alongside the school's 8-4-4.
        junior = YearGroup.objects.create(school=self.school_a, name="Grade 6", education_system="cbc",
                                          grading_scale="cbc")
        self.g6 = SchoolClass.objects.create(year_group=junior, name="6 Blue")
        self.pupil("Xen", self.east, {self.maths: 80})
        self.pupil("Yus", self.east, {self.maths: 50})
        self.pupil("Zed", self.g6, {self.maths: 95})

    def test_school_ranking_ranks_each_curriculum_separately_and_never_cbc(self):
        rows = self.rows(scope="school")
        self.assertEqual((rows["Xen"]["position"], rows["Xen"]["of"]), (1, 2))
        self.assertEqual(rows["Yus"]["position"], 2)
        self.assertIsNone(rows["Zed"]["position"])
        self.assertEqual(rows["Zed"]["section"], "CBC")


class UnsetCurriculumWithCbcLevelsTests(RankingFixture):
    """A school that never chose its curriculum but grades on CBC levels is treated as CBC: not ranked."""

    system, scale = "", "cbc4"

    def setUp(self):
        super().setUp()
        self.pupil("Amo", self.east, {self.maths: 90})
        self.pupil("Bea", self.east, {self.maths: 50})

    def test_not_ranked(self):
        for row in self.perf(scope="class", id=self.east.id)["students"]:
            self.assertIsNone(row["position"])
        response = self.admin.get("/api/exports/grades/", {"school_class": self.east.id, "term": self.t2.id})
        header = [c.value for c in openpyxl.load_workbook(io.BytesIO(response.content)).active[1]]
        self.assertNotIn("Position", header)
