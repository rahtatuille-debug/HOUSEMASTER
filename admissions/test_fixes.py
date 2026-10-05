"""
Fixes from the known-gaps list for admissions (B-1 to B-5). Test data is made up.
"""
import re
from datetime import timedelta
from io import StringIO

from django.core import mail
from django.core.management import call_command
from unittest import skipUnless

from django.db import connection
from django.test import TransactionTestCase
from django.utils import timezone

from activity.models import ActivityLog

from .models import Application
from .tests import AdmissionsFixture


def confirm_link(message):
    return re.search(r"/apply/confirm/(\S+)", message.body).group(1)


class ConfirmEmailTests(AdmissionsFixture):
    """B-1."""

    def test_an_application_waits_unconfirmed_and_one_email_goes_out(self):
        response = self.apply(confirm=False)
        self.assertEqual(response.status_code, 202, response.data)
        self.assertNotIn("reference", response.data)
        app = self.application()
        self.assertIsNone(app.confirmed_at)
        self.assertEqual(len(mail.outbox), 1)
        message = mail.outbox[0]
        self.assertEqual(message.to, ["priya@example.test"])
        self.assertIn("Alpha Academy", message.subject + message.body)  # says who it's from
        for detail in ("Zara", "Patel", "Hill Primary", "2015"):  # nothing about the child
            self.assertNotIn(detail, message.subject + message.body)
        token = confirm_link(message)
        self.assertNotIn(token, str(Application.objects.values().get()))  # only a hash is kept
        self.assertFalse(ActivityLog.objects.filter(action="admissions.applied").exists())

    def test_admins_see_only_confirmed_applications_and_unconfirmed_ones_separately(self):
        self.apply(confirm=False)
        app = self.application()
        self.assertEqual(self.admin.get("/api/admissions/applications/").data, [])
        self.assertEqual(self.admin.get(f"/api/admissions/applications/{app.id}/").status_code, 404)
        unconfirmed = self.admin.get("/api/admissions/applications/", {"unconfirmed": 1}).data
        self.assertEqual([a["id"] for a in unconfirmed], [app.id])
        summary = self.admin.get("/api/admissions/applications/summary/").data
        self.assertEqual((summary["counts"], summary["unconfirmed"]), ({}, 1))
        self.assertEqual(self.admin.post(f"/api/admissions/applications/{app.id}/enrol/",
                                         {"school_class": self.c7a.id}, format="json").status_code, 404)

    def test_confirming_releases_it_once(self):
        self.apply(confirm=False)
        token = confirm_link(mail.outbox[0])
        with self.captureOnCommitCallbacks(execute=True):
            response = self.public.post(f"/api/admissions/confirm/{token}/")
        self.assertEqual(response.status_code, 200, response.data)
        self.assertTrue(response.data["reference"].startswith("A"))
        self.assertIsNotNone(self.application().confirmed_at)
        self.assertIn(response.data["reference"], mail.outbox[-1].body)  # the receipt
        self.assertEqual(len(self.admin.get("/api/admissions/applications/").data), 1)
        self.assertTrue(ActivityLog.objects.filter(action="admissions.applied").exists())
        self.assertEqual(self.public.post(f"/api/admissions/confirm/{token}/").status_code, 400)  # single use
        self.assertEqual(self.public.post("/api/admissions/confirm/not-a-token/").status_code, 400)

    def test_links_expire_and_old_unconfirmed_applications_are_removed(self):
        self.apply(confirm=False)
        token = confirm_link(mail.outbox[0])
        Application.objects.update(confirm_expires_at=timezone.now() - timedelta(minutes=1))
        self.assertEqual(self.public.post(f"/api/admissions/confirm/{token}/").status_code, 400)
        self.apply(confirm=False, first_name="Omar", parent_email="o@example.test")  # still in time
        expired = self.application()
        out = StringIO()
        call_command("purge_applications", stdout=out)
        self.assertIn("1 unconfirmed", out.getvalue())
        self.assertEqual(Application.objects.count(), 2)  # dry run
        call_command("purge_applications", "--apply", stdout=StringIO())
        self.assertEqual(list(Application.objects.values_list("first_name", flat=True)), ["Omar"])
        self.assertNotIn(expired.first_name, out.getvalue())  # counts only


class DuplicateTests(AdmissionsFixture):
    """B-2."""

    def test_sending_the_same_child_twice_makes_one_application_and_one_email(self):
        first = self.apply(confirm=False)
        second = self.apply(confirm=False)
        self.assertEqual((first.status_code, second.status_code), (202, 202))
        self.assertEqual(first.data, second.data)  # nothing tells them apart
        self.assertEqual(Application.objects.count(), 1)
        self.assertEqual(len(mail.outbox), 1)

    def test_the_same_child_after_confirming_is_still_one_application(self):
        self.apply()
        self.apply(first_name=" zara ", last_name="PATEL", parent_email="Priya@Example.test")
        self.assertEqual(Application.objects.count(), 1)

    def test_a_different_child_or_a_closed_application_is_not_a_duplicate(self):
        self.apply()
        self.apply(first_name="Omar")  # a sibling
        self.assertEqual(Application.objects.count(), 2)
        Application.objects.filter(first_name="Zara").update(status="declined")
        self.apply()  # applying again after a decision
        self.assertEqual(Application.objects.filter(first_name="Zara").count(), 2)

    def test_an_expired_link_is_sent_again_rather_than_a_second_application(self):
        self.apply(confirm=False)
        Application.objects.update(confirm_expires_at=timezone.now() - timedelta(minutes=1))
        self.apply(confirm=False)
        self.assertEqual(Application.objects.count(), 1)
        self.assertEqual(len(mail.outbox), 2)
        self.assertEqual(self.public.post(f"/api/admissions/confirm/{confirm_link(mail.outbox[-1])}/").status_code, 200)

    def test_a_child_who_is_already_a_student_is_not_enrolled_again(self):
        from students.models import Student

        existing = Student.objects.create(school=self.school_a, first_name="Zara", last_name="Patel",
                                          date_of_birth="2015-03-02", school_class=self.c7a)
        self.apply()
        app = self.application()
        Application.objects.filter(pk=app.pk).update(status="offered")
        response = self.admin.post(f"/api/admissions/applications/{app.id}/enrol/", {"school_class": self.c7a.id},
                                   format="json")
        self.assertEqual(response.status_code, 400, response.data)
        self.assertEqual(int(response.data["existing_student"]), existing.id)
        self.assertEqual(Student.objects.filter(first_name="Zara").count(), 1)
        self.assertEqual(self.application().status, "offered")


@skipUnless(connection.vendor == "postgresql", "needs a database with concurrent writers (production's Postgres); "
                                               "SQLite's shared in-memory test database refuses the second writer")
class DuplicateRaceTests(TransactionTestCase):
    """B-2: two copies of the same form arriving at the same moment make one application."""

    def test_two_at_once(self):
        import threading

        from rest_framework.test import APIClient

        from students.models import School, YearGroup

        from .models import AdmissionsSettings

        school = School.objects.create(name="Race School")
        year = YearGroup.objects.create(school=school, name="Year 7")
        settings = AdmissionsSettings.objects.create(school=school, is_open=True)
        body = {"year_group": year.id, "first_name": "Ann", "last_name": "Lee", "date_of_birth": "2015-01-01",
                "parent_name": "Pat Lee", "parent_email": "pat@example.test", "parent_phone": "0700000000",
                "consent": True}
        barrier, codes = threading.Barrier(2), []

        def send():
            try:
                barrier.wait()
                codes.append(APIClient().post(f"/api/admissions/apply/{settings.token}/", body, format="json")
                             .status_code)
            finally:
                connection.close()

        threads = [threading.Thread(target=send) for _ in range(2)]
        for t in threads:
            t.start()
        for t in threads:
            t.join()
        self.assertEqual(codes, [202, 202])
        self.assertEqual(Application.objects.filter(school=school).count(), 1)


class ApplyLimitTests(AdmissionsFixture):
    """B-3."""

    def setUp(self):
        super().setUp()
        from django.core.cache import cache

        cache.clear()

    def test_a_school_or_cafe_sharing_one_address_is_not_blocked_early(self):
        codes = [self.apply(confirm=False, first_name=f"Kid{n}", parent_email=f"family{n}@example.test").status_code
                 for n in range(12)]  # the old limit was 10 an hour per address
        self.assertEqual(set(codes), {202})
        self.assertEqual(Application.objects.count(), 12)

    def test_one_email_address_is_limited_quietly(self):
        from unittest.mock import patch

        from .views import ApplyEmailThrottle

        with patch.dict(ApplyEmailThrottle.THROTTLE_RATES, {"admissions_apply_email": "2/hour"}):
            responses = [self.apply(confirm=False, first_name=f"Kid{n}") for n in range(3)]
        self.assertEqual([r.status_code for r in responses], [202, 202, 202])
        self.assertEqual(Application.objects.count(), 2)  # the third did nothing
        self.assertEqual(len(mail.outbox), 2)

    def test_the_answer_never_says_what_is_on_file(self):
        from unittest.mock import patch

        from .views import ApplyEmailThrottle

        new = self.apply(confirm=False)
        duplicate = self.apply(confirm=False)
        bot = self.apply(confirm=False, website="http://spam")
        with patch.dict(ApplyEmailThrottle.THROTTLE_RATES, {"admissions_apply_email": "1/hour"}):
            limited = self.apply(confirm=False, first_name="Omar")
        self.assertEqual({(r.status_code, str(r.data)) for r in (new, duplicate, bot, limited)},
                         {(new.status_code, str(new.data))})

    def test_the_address_limit_still_applies(self):
        from unittest.mock import patch

        from .views import ApplyThrottle

        with patch.dict(ApplyThrottle.THROTTLE_RATES, {"admissions_apply": "2/hour"}):
            codes = [self.apply(confirm=False, first_name=f"K{n}", parent_email=f"f{n}@example.test").status_code
                     for n in range(3)]
        self.assertEqual(codes, [202, 202, 429])


class FamilyDataTests(AdmissionsFixture):
    """B-4."""

    def enrolled(self):
        self.apply(medical_notes="Asthma: inhaler in bag")
        app = self.application()
        Application.objects.filter(pk=app.pk).update(status="offered", staff_notes="Met the head")
        response = self.admin.post(f"/api/admissions/applications/{app.id}/enrol/", {"school_class": self.c7a.id},
                                   format="json")
        return app, response.data["student"]

    def test_the_family_export_includes_their_applications(self):
        from openpyxl import load_workbook
        from io import BytesIO

        app, student_id = self.enrolled()
        # An earlier application for the same child, declined.
        old = Application.objects.create(school=self.school_a, first_name="Zara", last_name="Patel",
                                         date_of_birth="2015-03-02", parent_name="Priya Patel",
                                         parent_email="priya@example.test", status="declined",
                                         confirmed_at=timezone.now())
        data = self.admin.get(f"/api/students/{student_id}/data-export/", {"format": "json"}).data
        rows = data["admissions_applications"]
        self.assertEqual({r["reference"] for r in rows}, {app.reference, old.reference})
        self.assertIn("Asthma: inhaler in bag", str(rows))
        response = self.admin.get(f"/api/students/{student_id}/data-export/")
        sheet = load_workbook(BytesIO(response.content))["Applications"]
        self.assertIn(app.reference, str([c.value for row in sheet.iter_rows() for c in row]))

    def test_removing_a_familys_data_removes_their_applications(self):
        app, student_id = self.enrolled()
        Application.objects.create(school=self.school_b, first_name="Zara", last_name="Patel",
                                   date_of_birth="2015-03-02", parent_name="X", parent_email="x@example.test")
        response = self.admin.post(f"/api/students/{student_id}/remove-personal-data/",
                                   {"confirm_name": "Zara Patel"}, format="json")
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["applications_deleted"], 1)
        self.assertFalse(Application.objects.filter(school=self.school_a).exists())
        self.assertTrue(Application.objects.filter(school=self.school_b).exists())  # another school's is untouched

    def test_retention_is_off_by_default_and_a_dry_run_by_default(self):
        self.apply()
        self.apply(first_name="Omar", parent_email="o@example.test")
        old = timezone.now() - timedelta(days=400)
        Application.objects.filter(first_name="Zara").update(status="declined", updated_at=old)
        Application.objects.filter(first_name="Omar").update(updated_at=old)  # still open: never removed
        call_command("purge_applications", "--apply", stdout=StringIO())
        self.assertEqual(Application.objects.count(), 2)  # off by default
        settings = self.admin.patch("/api/admissions/settings/", {"retention_days": 365}, format="json")
        self.assertEqual(settings.data["retention_days"], 365)
        out = StringIO()
        call_command("purge_applications", stdout=out)
        self.assertIn("1 closed", out.getvalue())
        self.assertEqual(Application.objects.count(), 2)
        call_command("purge_applications", "--apply", stdout=StringIO())
        self.assertEqual(list(Application.objects.values_list("first_name", flat=True)), ["Omar"])
        self.assertNotIn("Zara", out.getvalue())


class FormChecksTests(AdmissionsFixture):
    """B-5: dates of birth, phones, the age banner and admission numbers."""

    def test_a_date_of_birth_in_the_future_or_implausibly_long_ago_is_refused(self):
        tomorrow = (timezone.localdate() + timedelta(days=1)).isoformat()
        response = self.apply(confirm=False, date_of_birth=tomorrow)
        self.assertEqual(response.status_code, 400)
        self.assertIn("date_of_birth", response.data)
        self.assertEqual(self.apply(confirm=False, date_of_birth="1960-05-01").status_code, 400)
        self.assertFalse(Application.objects.exists())

    def test_phones_are_checked_leniently(self):
        for good in ("0712 345 678", "+254 712 345 678", "(020) 123-4567", "+44 20 7946 0958", "0712345678"):
            self.assertNotEqual(self.apply(confirm=False, parent_phone=good, first_name=f"K{len(good)}x{good[-2:]}")
                                .status_code, 400, good)
        for bad in ("12", "call me", "07123456789012345678", "+-+"):
            response = self.apply(confirm=False, parent_phone=bad, first_name="Bad")
            self.assertEqual(response.status_code, 400, bad)
            self.assertIn("parent_phone", response.data)

    def test_an_age_far_from_the_year_group_is_flagged_not_blocked(self):
        from students.models import Student

        for n in range(3):  # Year 7 is full of eleven-year-olds
            Student.objects.create(school=self.school_a, first_name=f"S{n}", last_name="Y",
                                   date_of_birth=timezone.localdate() - timedelta(days=int(11.5 * 365)),
                                   school_class=self.c7a)
        young = (timezone.localdate() - timedelta(days=5 * 365 + 100)).isoformat()
        self.assertEqual(self.apply(date_of_birth=young).status_code, 200)  # not blocked
        row = self.admin.get("/api/admissions/applications/").data[0]
        self.assertIn("most students in Year 7 are 11", row["age_note"])
        self.apply(first_name="Omar", parent_email="o@example.test",
                   date_of_birth=(timezone.localdate() - timedelta(days=int(11.2 * 365))).isoformat())
        omar = next(r for r in self.admin.get("/api/admissions/applications/").data if r["first_name"] == "Omar")
        self.assertEqual(omar["age_note"], "")

    def enrol(self, first_name="Zara", **extra):
        self.apply(first_name=first_name, parent_email=f"{first_name.lower()}@example.test", **extra)
        app = Application.objects.get(first_name=first_name)
        Application.objects.filter(pk=app.pk).update(status="offered")
        return self.admin.post(f"/api/admissions/applications/{app.id}/enrol/", {"school_class": self.c7a.id},
                               format="json")

    def test_enrolling_gives_the_next_admission_number_with_the_prefix(self):
        from students.models import Student

        Student.objects.create(school=self.school_a, first_name="Old", last_name="One", external_id="ADM-1")
        self.admin.patch("/api/admissions/settings/", {"number_prefix": "ADM-"}, format="json")
        first = Student.objects.get(pk=self.enrol().data["student"])
        second = Student.objects.get(pk=self.enrol("Omar").data["student"])
        self.assertEqual((first.external_id, second.external_id), ("ADM-2", "ADM-3"))  # ADM-1 was taken
        self.assertEqual(Student.objects.get(first_name="Old").external_id, "ADM-1")  # never renumbered
        self.assertIn("ADM-2", self.admin.get(f"/api/admissions/applications/{Application.objects.get(first_name='Zara').id}/")
                      .data["student_number"])

    def test_admission_numbers_are_unique_per_school(self):
        from django.db import IntegrityError, transaction

        from students.models import Student

        Student.objects.create(school=self.school_a, first_name="A", last_name="A", external_id="7")
        Student.objects.create(school=self.school_b, first_name="B", last_name="B", external_id="7")  # fine
        Student.objects.create(school=self.school_a, first_name="C", last_name="C", external_id="")
        Student.objects.create(school=self.school_a, first_name="D", last_name="D", external_id="")  # blanks fine
        with self.assertRaises(IntegrityError), transaction.atomic():
            Student.objects.create(school=self.school_a, first_name="E", last_name="E", external_id="7")
        response = self.admin.post("/api/students/", {"first_name": "F", "last_name": "F", "external_id": "7"},
                                   format="json")
        self.assertEqual(response.status_code, 400, response.data)
        self.assertIn("external_id", response.data)


@skipUnless(connection.vendor == "postgresql", "needs a database with concurrent writers (production's Postgres)")
class AdmissionNumberRaceTests(TransactionTestCase):
    """B-5: two enrolments at the same moment get different numbers."""

    def test_two_at_once(self):
        import threading

        from rest_framework.test import APIClient

        from accounts.models import Profile
        from django.contrib.auth.models import User
        from students.models import School, SchoolClass, Student, YearGroup

        from .models import AdmissionsSettings

        school = School.objects.create(name="Race School")
        klass = SchoolClass.objects.create(year_group=YearGroup.objects.create(school=school, name="Y7"), name="7A")
        AdmissionsSettings.objects.create(school=school, number_prefix="R")
        admin = User.objects.create_user(username="race@example.test", email="race@example.test", password="x")
        Profile.objects.create(user=admin, school=school, role="admin")
        apps = [Application.objects.create(school=school, first_name=f"Kid{n}", last_name="Lee", status="offered",
                                           parent_name="P", parent_email=f"p{n}@example.test",
                                           confirmed_at=timezone.now()) for n in range(2)]
        barrier, codes = threading.Barrier(2), []

        def enrol(app):
            try:
                client = APIClient()
                client.force_authenticate(admin)
                barrier.wait()
                codes.append(client.post(f"/api/admissions/applications/{app.id}/enrol/", {"school_class": klass.id},
                                         format="json").status_code)
            finally:
                connection.close()

        threads = [threading.Thread(target=enrol, args=(a,)) for a in apps]
        for t in threads:
            t.start()
        for t in threads:
            t.join()
        self.assertEqual(codes, [200, 200])
        self.assertEqual(sorted(Student.objects.filter(school=school).values_list("external_id", flat=True)),
                         ["R1", "R2"])
