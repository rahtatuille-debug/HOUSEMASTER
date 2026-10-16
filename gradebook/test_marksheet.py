"""Mark sheets: download a class's sheet, fill in marks in Excel, check it, then import it."""
from io import BytesIO

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile

from accounts.tests import SchoolScopedAPITestCase
from activity.models import ActivityLog
from gradebook.models import AssessmentType, Grade, StudentSubject, Subject, Term
from students.models import SchoolClass, Student, YearGroup

TEMPLATE = "/api/marksheets/template/"
IMPORT = "/api/marksheets/import/"


class MarkSheetTests(SchoolScopedAPITestCase):
    def setUp(self):
        super().setUp()
        school = self.school_a
        year = YearGroup.objects.create(school=school, name="Year 7")
        self.c7a = SchoolClass.objects.create(year_group=year, name="7A")
        self.c7b = SchoolClass.objects.create(year_group=year, name="7B")
        self.term = Term.objects.create(school=school, name="Term 1")
        self.maths = Subject.objects.create(school=school, name="Maths")
        self.french = Subject.objects.create(school=school, name="French", is_elective=True)
        self.cat = AssessmentType.objects.create(school=school, name="CAT", weight=30)

        def kid(first, adm, klass=self.c7a, active=True):
            return Student.objects.create(school=school, first_name=first, last_name="X", external_id=adm,
                                          school_class=klass, is_active=active)

        self.ann, self.ben, self.cy = kid("Ann", "ADM/1"), kid("Ben", "ADM/2"), kid("Cy", "ADM/3")
        self.dee = kid("Dee", "ADM/4", klass=self.c7b)
        kid("Gone", "ADM/5", active=False)
        StudentSubject.objects.create(student=self.ann, subject=self.french)  # only Ann takes French
        self.assign(self.user_a, self.c7a, self.maths)
        self.teacher = self.client_a
        self.params = {"school_class": self.c7a.id, "subject": self.maths.id, "term": self.term.id,
                       "assessment_type": self.cat.id, "max_score": "50"}

    def sheet(self, marks, params=None):
        """The downloaded sheet with marks typed in: {admission number: mark}; extra rows appended."""
        response = self.teacher.get(TEMPLATE, params or self.params)
        self.assertEqual(response.status_code, 200, response.content[:200])
        wb = openpyxl.load_workbook(BytesIO(response.content))
        ws = wb["Marks"]
        listed = {}
        for row in ws.iter_rows(min_row=6):
            if row[0].value:
                listed[row[0].value] = row
        for adm, mark in marks.items():
            if adm in listed:
                listed[adm][2].value = mark
            else:
                ws.append([adm, "", mark])
        out = BytesIO()
        wb.save(out)
        return SimpleUploadedFile("marks.xlsx", out.getvalue())

    def upload(self, file, commit=False, client=None, **params):
        return (client or self.teacher).post(IMPORT, {**self.params, **params, "file": file, "commit": "true" if commit else "false"},
                                              format="multipart")

    def test_the_sheet_lists_the_class_students_who_take_the_subject(self):
        response = self.teacher.get(TEMPLATE, self.params)
        ws = openpyxl.load_workbook(BytesIO(response.content))["Marks"]
        self.assertIn("Maths · 7A · Term 1", ws["A1"].value)
        self.assertIn("CAT · Out of 50", ws["A2"].value)
        self.assertEqual([c.value for c in ws[5]], ["Admission no.", "Student", "Mark"])
        self.assertEqual([r[0] for r in ws.iter_rows(min_row=6, values_only=True)], ["ADM/1", "ADM/2", "ADM/3"])
        # An elective: only the students who chose it.
        self.assign(self.user_a, self.c7a, self.french)
        response = self.teacher.get(TEMPLATE, {**self.params, "subject": self.french.id})
        ws = openpyxl.load_workbook(BytesIO(response.content))["Marks"]
        self.assertEqual([r[0] for r in ws.iter_rows(min_row=6, values_only=True)], ["ADM/1"])

    def test_preview_checks_every_row_and_saves_nothing(self):
        file = self.sheet({"ADM/1": 42, "ADM/2": "ABS", "ADM/4": 30, "ADM/9": 10, "ADM/1 ": 20})
        data = self.upload(file).data
        self.assertEqual(Grade.objects.count(), 0)
        rows = {(r["admission_number"], r["status"]): r["message"] for r in data["rows"]}
        self.assertEqual(rows[("ADM/1", "ok")], "")
        self.assertEqual(rows[("ADM/2", "skipped")], "No mark: skipped.")
        self.assertEqual(rows[("ADM/4", "error")], "Not in 7A.")
        self.assertEqual(rows[("ADM/9", "error")], "No student has this admission number.")
        self.assertEqual(rows[("ADM/1", "error")], "This student is in the sheet twice.")
        # Cy's row is still there with no mark: skipped too.
        self.assertEqual(data["counts"], {"ready": 1, "skipped": 2, "errors": 3, "missing": 0})
        self.assertEqual(data["imported"], 0)

    def test_a_student_left_off_the_sheet_is_listed_as_missing(self):
        response = self.teacher.get(TEMPLATE, self.params)
        wb = openpyxl.load_workbook(BytesIO(response.content))
        ws = wb["Marks"]
        ws.delete_rows(8)  # Cy's row
        ws["C6"], ws["C7"] = 40, 45
        out = BytesIO()
        wb.save(out)
        data = self.upload(SimpleUploadedFile("marks.xlsx", out.getvalue())).data
        self.assertEqual((data["counts"]["ready"], data["missing"]), (2, ["Cy X"]))

    def test_bad_marks(self):
        data = self.upload(self.sheet({"ADM/1": 51, "ADM/2": "forty", "ADM/3": -1})).data
        messages = [r["message"] for r in data["rows"]]
        self.assertEqual(messages, ["The mark must be between 0 and 50.", "The mark isn't a number.",
                                    "The mark must be between 0 and 50."])

    def test_import_saves_the_marks(self):
        data = self.upload(self.sheet({"ADM/1": 42, "ADM/2": "37.5", "ADM/3": ""}), commit=True).data
        self.assertEqual(data["imported"], 2)
        marks = {g.student_id: g for g in Grade.objects.all()}
        self.assertEqual(set(marks), {self.ann.id, self.ben.id})
        g = marks[self.ben.id]
        self.assertEqual((float(g.score), float(g.max_score), g.assessment_type, g.subject, g.term),
                         (37.5, 50.0, self.cat, self.maths, self.term))
        log = ActivityLog.objects.get(action="grade.imported")
        self.assertIn("Imported 2 Maths marks for 7A", log.summary)

    def test_marks_without_an_assessment_type(self):
        params = {**self.params, "assessment_type": "", "max_score": ""}
        data = self.upload(self.sheet({"ADM/1": 90}, params=params), commit=True, assessment_type="", max_score="").data
        self.assertEqual(data["imported"], 1)
        g = Grade.objects.get()
        self.assertEqual((g.assessment_type, float(g.max_score)), (None, 100.0))

    def test_only_a_teacher_of_the_class_and_subject(self):
        self.assign(self.user_a, self.c7b, self.french)
        for params in ({**self.params, "school_class": self.c7b.id}, {**self.params, "subject": self.french.id}):
            self.assertEqual(self.teacher.get(TEMPLATE, params).status_code, 403)
        file = SimpleUploadedFile("marks.xlsx", b"x")
        self.assertEqual(self.upload(file, school_class=self.c7b.id).status_code, 403)
        # Admins may for any class.
        admin = self.authed_client(self.admin_a)
        self.assertEqual(admin.get(TEMPLATE, {**self.params, "school_class": self.c7b.id}).status_code, 200)

    def test_another_schools_class_is_not_found(self):
        other = SchoolClass.objects.create(year_group=YearGroup.objects.create(school=self.school_b, name="Y"), name="Z")
        self.make_admin(self.user_b)
        self.assertEqual(self.client_b.get(TEMPLATE, self.params).status_code, 404)  # school_a's ids
        self.assertEqual(self.teacher.get(TEMPLATE, {**self.params, "school_class": other.id}).status_code, 404)

    def test_locked_term(self):
        file = self.sheet({"ADM/1": 42})
        self.term.is_locked = True
        self.term.save()
        response = self.upload(file, commit=True)
        self.assertEqual(response.status_code, 400)
        self.assertIn("locked", str(response.data))
        self.assertEqual(Grade.objects.count(), 0)

    def test_unreadable_files(self):
        response = self.upload(SimpleUploadedFile("marks.xlsx", b"not a spreadsheet"))
        self.assertEqual(response.status_code, 400)
        self.assertIn("isn't an Excel file", str(response.data["file"]))
        wb = openpyxl.Workbook()
        wb.active.append(["Name", "Score"])
        out = BytesIO()
        wb.save(out)
        response = self.upload(SimpleUploadedFile("other.xlsx", out.getvalue()))
        self.assertIn("Admission no.", str(response.data["file"]))
        self.assertEqual(self.upload(self.sheet({}), max_score="0").status_code, 400)

    def test_a_zip_bomb_is_refused_before_it_is_unpacked(self):
        from housemaster.test_uploads import raw_xlsx

        bomb = raw_xlsx(shared_text_mb=30)  # under 2 MB, unpacks to about 30 MB
        response = self.upload(SimpleUploadedFile("marks.xlsx", bomb), commit=True)
        self.assertEqual(response.status_code, 400)
        self.assertIn("unpacks to far more data", str(response.data["file"]))
        self.assertEqual(Grade.objects.count(), 0)

    def test_a_sheet_of_empty_rows_is_refused(self):
        from housemaster.test_uploads import raw_xlsx

        sheet = raw_xlsx(header=("Admission no.", "Student", "Mark"), empty_rows=25_000)
        response = self.upload(SimpleUploadedFile("marks.xlsx", sheet))
        self.assertEqual(response.status_code, 400)
        self.assertIn("most of them empty", str(response.data["file"]))
