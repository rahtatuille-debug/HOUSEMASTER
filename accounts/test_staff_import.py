"""Importing staff from Excel: invites with their classes, preview, emails and acceptance."""
from io import BytesIO

import openpyxl
from django.core import mail
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from rest_framework.test import APIClient

from activity.models import ActivityLog
from gradebook.models import Subject
from students.models import SchoolClass, YearGroup

from .models import Invite, Profile, TeachingAssignment
from .tests import SchoolScopedAPITestCase

HEADER = ["name", "email", "role", "class_teacher_of", "teaches"]


def sheet(*rows, header=HEADER):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Staff"
    ws.append(header)
    for row in rows:
        ws.append(list(row))
    out = BytesIO()
    wb.save(out)
    return SimpleUploadedFile("staff.xlsx", out.getvalue(),
                              content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend", NOTIFICATIONS_IN_BACKGROUND=False)
class StaffImportTests(SchoolScopedAPITestCase):
    def setUp(self):
        super().setUp()
        self.admin = self.authed_client(self.admin_a)
        g7 = YearGroup.objects.create(school=self.school_a, name="Grade 7")
        g8 = YearGroup.objects.create(school=self.school_a, name="Grade 8")
        self.c7e = SchoolClass.objects.create(year_group=g7, name="7 East")
        self.c7w = SchoolClass.objects.create(year_group=g7, name="7 West")
        self.east7 = SchoolClass.objects.create(year_group=g7, name="East")
        self.east8 = SchoolClass.objects.create(year_group=g8, name="East")
        self.maths = Subject.objects.create(school=self.school_a, name="Mathematics")
        self.english = Subject.objects.create(school=self.school_a, name="English")
        # Another school with the same names must never be matched.
        other = YearGroup.objects.create(school=self.school_b, name="Grade 7")
        SchoolClass.objects.create(year_group=other, name="7 South")

    def post(self, upload, **extra):
        return self.admin.post("/api/import/staff/", {"file": upload, **extra}, format="multipart")

    def test_preview_saves_nothing_and_shows_what_would_happen(self):
        response = self.post(sheet(["Mary Njeri", "Mary@Alpha.test", "", "7 East", "7 West: Mathematics, English"]))
        self.assertEqual(response.status_code, 200)
        person = response.data["people"][0]
        self.assertEqual((person["email"], person["role"]), ("mary@alpha.test", "teacher"))
        self.assertEqual(person["assignments"],
                         ["7 East (all subjects)", "7 West (Mathematics)", "7 West (English)"])
        self.assertNotIn("token", person)
        self.assertFalse(Invite.objects.filter(email="mary@alpha.test").exists())

    def test_import_creates_invites_that_assign_classes_on_acceptance(self):
        response = self.post(sheet(["Mary Njeri", "mary@alpha.test", "teacher", "7 East", "7 West: Mathematics"],
                                   ["Pat Head", "pat@alpha.test", "admin", "", ""]), commit="true")
        self.assertEqual(len(response.data["people"]), 2)
        invite = Invite.objects.get(email="mary@alpha.test")
        self.assertEqual(response.data["people"][0]["token"], invite.token)
        self.assertEqual(Invite.objects.get(email="pat@alpha.test").role, "admin")
        self.assertTrue(ActivityLog.objects.filter(action="staff.imported").exists())

        accepted = APIClient().post("/api/invites/accept/", {"token": invite.token, "password": "a-long-Password-123",
                                                             "accept_privacy": True})
        self.assertEqual(accepted.status_code, 201)
        profile = Profile.objects.get(user__email="mary@alpha.test")
        self.assertEqual(
            set(TeachingAssignment.objects.filter(teacher=profile).values_list("school_class__name", "subject__name")),
            {("7 East", None), ("7 West", "Mathematics")},
        )

    def test_row_errors_are_reported_and_good_rows_still_import(self):
        response = self.post(sheet(
            ["", "no.name@alpha.test", "", "", ""],
            ["Bad Email", "not-an-email", "", "", ""],
            ["Bad Role", "role@alpha.test", "principal", "", ""],
            ["No Class", "nc@alpha.test", "", "9 North", ""],
            ["No Subject", "ns@alpha.test", "", "", "7 East: Art"],
            ["Bad Format", "bf@alpha.test", "", "", "7 East Mathematics"],
            ["Ambiguous", "amb@alpha.test", "", "East", ""],
            ["Other School", "os@alpha.test", "", "7 South", ""],
            ["Good One", "good@alpha.test", "", "Grade 8/East", ""],
        ), commit="true")
        self.assertEqual(len(response.data["errors"]), 8)
        self.assertEqual([p["email"] for p in response.data["people"]], ["good@alpha.test"])
        self.assertEqual(response.data["people"][0]["assignments"], ["East (all subjects)"])
        self.assertEqual(Invite.objects.get(email="good@alpha.test").assignments,
                         [{"school_class": self.east8.id, "subject": None}])

    def test_existing_people_and_duplicates_are_skipped(self):
        Invite.objects.create(school=self.school_a, email="pending@alpha.test", name="P", invited_by=self.admin_a)
        response = self.post(sheet(["Already", "teacher.a@alpha.test", "", "", ""],
                                   ["Pending", "pending@alpha.test", "", "", ""],
                                   ["Twice", "twice@alpha.test", "", "", ""],
                                   ["Twice again", "twice@alpha.test", "", "", ""]), commit="true")
        self.assertEqual([s["email"] for s in response.data["skipped"]], ["teacher.a@alpha.test", "pending@alpha.test"])
        self.assertEqual(len(response.data["errors"]), 1)
        self.assertEqual(Invite.objects.filter(email="twice@alpha.test").count(), 1)

    def test_invite_links_can_be_emailed(self):
        with self.captureOnCommitCallbacks(execute=True):
            self.post(sheet(["Mary Njeri", "mary@alpha.test", "", "", ""]), commit="true", send_emails="true")
        self.assertEqual(len(mail.outbox), 1)
        token = Invite.objects.get(email="mary@alpha.test").token
        self.assertIn(f"/invite/{token}", mail.outbox[0].body)
        self.assertEqual(mail.outbox[0].to, ["mary@alpha.test"])

    def test_no_emails_unless_asked(self):
        with self.captureOnCommitCallbacks(execute=True):
            self.post(sheet(["Mary Njeri", "mary@alpha.test", "", "", ""]), commit="true")
        self.assertEqual(mail.outbox, [])

    def test_only_admins_can_import_or_get_the_template(self):
        self.assertEqual(self.client_a.post("/api/import/staff/", {"file": sheet()}, format="multipart").status_code, 403)
        self.assertEqual(self.client_a.get("/api/import/staff-template/").status_code, 403)
        template = self.admin.get("/api/import/staff-template/")
        ws = openpyxl.load_workbook(BytesIO(template.content)).active
        self.assertEqual([c.value for c in ws[1]], HEADER)

    def test_not_an_excel_file(self):
        bad = SimpleUploadedFile("staff.xlsx", b"hello", content_type="application/octet-stream")
        self.assertEqual(self.post(bad).status_code, 400)
