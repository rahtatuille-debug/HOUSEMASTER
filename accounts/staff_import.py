"""
Import staff from an Excel sheet: one invite per person, with the classes
and subjects they teach assigned automatically when they accept.

Columns (first row is the header, order doesn't matter):
  name, email, role, class_teacher_of, teaches

* role: teacher (the default) or admin.
* class_teacher_of: classes they teach every subject in, separated by ";".
* teaches: "class: subject, subject" entries separated by ";",
  e.g. "7 West: Mathematics; 8 East: Mathematics, English".

A class can be written "Grade 7/East" when two year groups share a class
name. Classes and subjects must already exist. People who already have an
account or a pending invite at the school are skipped. Preview and import
run the same code: a preview rolls everything back.

Each invite counts against the same limits as a single invite (per admin
per hour, per recipient per day; accounts.throttles.reserve_invite_email).
Rows over a limit are listed as `deferred` and nothing is created for
them; running the same sheet again later invites them, since the people
already invited are skipped.
"""
from collections import Counter

import openpyxl
from django.core.exceptions import ValidationError as DjangoValidationError
from django.core.validators import validate_email
from django.db import transaction

from gradebook.models import Subject
from reporting.spreadsheets import append_row
from students.importer import MAX_ERRORS, WorkbookError, _rows, _text
from students.models import SchoolClass

from .models import Invite, Profile
from .serializers import email_in_use_at_school
from .throttles import reserve_invite_email

SHEET = "Staff"
COLUMNS = ["name", "email", "role", "class_teacher_of", "teaches"]
EXAMPLES = [
    ["Mary Njeri", "m.njeri@yourschool.ac.ke", "teacher", "7 East", "7 West: Mathematics; 8 East: Mathematics"],
    ["Peter Otieno", "p.otieno@yourschool.ac.ke", "admin", "", ""],
]
ROLES = {"teacher": Profile.Role.TEACHER, "admin": Profile.Role.ADMIN, "": Profile.Role.TEACHER}


def staff_template(school=None):
    """
    The staff sheet with example rows. For a school, the examples use its own
    classes and subjects, and a "Classes and subjects" sheet lists the names
    the import will recognise.
    """
    examples = [list(r) for r in EXAMPLES]
    classes, subjects = [], []
    if school is not None:
        classes = [c.name for c in SchoolClass.objects.filter(year_group__school=school)
                   .order_by("year_group__order", "year_group__name", "name")]
        subjects = list(Subject.objects.filter(school=school).order_by("name").values_list("name", flat=True))
        if classes and subjects:
            other = classes[1] if len(classes) > 1 else classes[0]
            examples[0][3] = classes[0]
            examples[0][4] = f"{other}: {subjects[0]}" + (f"; {classes[0]}: {subjects[1]}" if len(subjects) > 1 else "")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    append_row(ws, COLUMNS)
    for row in examples:
        append_row(ws, row)
    for letter, width in zip("ABCDE", (22, 30, 10, 22, 50)):
        ws.column_dimensions[letter].width = width
    if classes or subjects:
        lists = wb.create_sheet("Classes and subjects")
        append_row(lists, ["class", "subject"])
        for i in range(max(len(classes), len(subjects))):
            append_row(lists, [classes[i] if i < len(classes) else None, subjects[i] if i < len(subjects) else None])
        lists.column_dimensions["A"].width = lists.column_dimensions["B"].width = 28
    return wb


class _Rollback(Exception):
    pass


class _StaffImport:
    def __init__(self, school, invited_by, commit):
        self.school = school
        self.invited_by = invited_by
        self.commit = commit
        self.pending = Counter()  # a preview's invite counts (it must not use up the real limits)
        self.errors, self.people, self.skipped, self.deferred = [], [], [], []
        self.classes = list(SchoolClass.objects.filter(year_group__school=school).select_related("year_group"))
        # Keyed by curriculum too: a school running two can have "Mathematics" in each.
        self.subjects = {(s.name.lower(), s.education_system): s for s in Subject.objects.filter(school=school)}
        self.seen_emails = set()

    def error(self, row, message):
        if len(self.errors) < MAX_ERRORS:
            self.errors.append({"sheet": SHEET, "row": row, "message": message})

    def find_class(self, text):
        text = " ".join(text.split())
        if "/" in text:
            year, _, name = (part.strip().lower() for part in text.partition("/"))
            matches = [c for c in self.classes if c.name.lower() == name and c.year_group.name.lower() == year]
        else:
            matches = [c for c in self.classes if c.name.lower() == text.lower()]
        if len(matches) == 1:
            return matches[0]
        if matches:
            raise ValueError(f'"{text}" matches more than one class; write it as "Year group/Class"')
        raise ValueError(f'no class called "{text}" (add it in Setup first)')

    def find_subject(self, text, school_class):
        from students.presets import section_for, subject_key

        system = section_for(school_class.year_group, self.school)[0]
        subject = self.subjects.get((" ".join(text.split()).lower(), subject_key(self.school, system)))
        if subject is None:
            raise ValueError(f'no subject called "{text}" in {school_class.name}\'s curriculum (add it in Setup first)')
        return subject

    def assignments(self, row):
        """[(class, subject or None)] from the class_teacher_of and teaches columns."""
        result = []
        for part in _text(row.get("class_teacher_of")).split(";"):
            if part.strip():
                result.append((self.find_class(part), None))
        for entry in _text(row.get("teaches")).split(";"):
            if not entry.strip():
                continue
            class_text, colon, subjects = entry.partition(":")
            if not colon or not subjects.strip():
                raise ValueError(f'"{entry.strip()}" should look like "7 West: Mathematics"')
            school_class = self.find_class(class_text)
            for subject_text in subjects.split(","):
                if subject_text.strip():
                    result.append((school_class, self.find_subject(subject_text, school_class)))
        return list(dict.fromkeys(result))

    def run(self, ws):
        for number, row in _rows(ws):
            name, email = " ".join(_text(row.get("name")).split()), _text(row.get("email")).lower()
            role_text = _text(row.get("role")).lower()
            if not name or not email:
                self.error(number, "name and email are both required")
                continue
            try:
                validate_email(email)
            except DjangoValidationError:
                self.error(number, f'"{email}" isn\'t an email address')
                continue
            if role_text not in ROLES:
                self.error(number, f'role "{role_text}" must be teacher or admin')
                continue
            if email in self.seen_emails:
                self.error(number, f"{email} appears more than once in the sheet")
                continue
            self.seen_emails.add(email)
            try:
                assignments = self.assignments(row)
            except ValueError as exc:
                self.error(number, str(exc))
                continue
            if email_in_use_at_school(email, self.school) or Invite.objects.filter(
                    school=self.school, email__iexact=email, accepted_at__isnull=True).exists():
                self.skipped.append({"row": number, "name": name, "email": email,
                                     "reason": "already has an account or a pending invite"})
                continue
            refusal = reserve_invite_email(self.invited_by, email, record=self.commit, pending=self.pending)
            if refusal:
                self.deferred.append({"row": number, "name": name, "reason": refusal})
                continue
            invite = Invite.objects.create(
                school=self.school, name=name, email=email, role=ROLES[role_text], invited_by=self.invited_by,
                assignments=[{"school_class": c.id, "subject": s.id if s else None} for c, s in assignments],
            )
            self.people.append({
                "row": number, "name": name, "email": email, "role": invite.role, "invite_id": invite.id,
                "token": invite.token,
                "assignments": [f"{c.name} ({s.name if s else 'all subjects'})" for c, s in assignments],
            })


def import_staff(upload, school, invited_by, commit=False):
    try:
        wb = openpyxl.load_workbook(upload, read_only=True, data_only=True)
    except Exception:
        raise WorkbookError("That file isn't an Excel workbook (.xlsx).")
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    job = _StaffImport(school, invited_by, commit)
    try:
        with transaction.atomic():
            job.run(ws)
            if not commit:
                raise _Rollback
    except _Rollback:
        for person in job.people:  # nothing was saved, so there are no links yet
            person.pop("token")
            person.pop("invite_id")
    return {"committed": commit, "people": job.people, "skipped": job.skipped, "deferred": job.deferred,
            "errors": job.errors}
