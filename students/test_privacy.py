"""Data protection: the privacy notice at signup, and a family's data export and removal."""
from datetime import date
from io import BytesIO

import openpyxl
from django.contrib.auth.models import User
from rest_framework.test import APIClient

from accounts.models import Invite, Profile
from accounts.tests import SchoolScopedAPITestCase
from activity.models import ActivityLog
from attendance.models import AttendanceRecord
from gradebook.models import Grade, Subject, Term
from guardians.models import Guardian, GuardianInvite
from messaging.models import Conversation, ConversationParticipant, Message
from reporting.models import StudentReport

from .models import SchoolClass, Student, YearGroup


class PrivacyNoticeTests(SchoolScopedAPITestCase):
    def setUp(self):
        super().setUp()
        self.school_a.privacy_contact = "privacy@alpha.test"
        self.school_a.save()
        self.anon = APIClient()

    def test_staff_must_accept_the_notice_and_it_is_recorded(self):
        invite = Invite.objects.create(school=self.school_a, email="new@alpha.test", name="New Teacher",
                                       invited_by=self.admin_a)
        preview = self.anon.get(f"/api/invites/preview/{invite.token}/").data
        self.assertEqual(preview["privacy_contact"], "privacy@alpha.test")
        body = {"token": invite.token, "password": "a-long-Password-123"}
        self.assertEqual(self.anon.post("/api/invites/accept/", body).status_code, 400)
        self.assertEqual(self.anon.post("/api/invites/accept/", {**body, "accept_privacy": False}).status_code, 400)
        self.assertEqual(self.anon.post("/api/invites/accept/", {**body, "accept_privacy": True}).status_code, 201)
        self.assertIsNotNone(Profile.objects.get(user__email="new@alpha.test").privacy_accepted_at)

    def test_parents_must_accept_the_notice_and_it_is_recorded(self):
        kid = Student.objects.create(school=self.school_a, first_name="Kid", last_name="One")
        invite = GuardianInvite.objects.create(school=self.school_a, email="mum@alpha.test", name="Mum",
                                               invited_by=self.admin_a)
        invite.students.add(kid)
        self.assertEqual(self.anon.get(f"/api/guardian-invites/preview/{invite.token}/").data["privacy_contact"],
                         "privacy@alpha.test")
        body = {"token": invite.token, "password": "a-long-Password-123"}
        self.assertEqual(self.anon.post("/api/guardian-invites/accept/", body).status_code, 400)
        self.assertEqual(self.anon.post("/api/guardian-invites/accept/", {**body, "accept_privacy": True}).status_code, 201)
        self.assertIsNotNone(Guardian.objects.get(user__email="mum@alpha.test").privacy_accepted_at)

    def test_admin_sets_the_privacy_contact(self):
        admin = self.authed_client(self.admin_a)
        response = admin.patch(f"/api/schools/{self.school_a.id}/", {"privacy_contact": "dpo@alpha.test"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["privacy_contact"], "dpo@alpha.test")


class FamilyDataFixture(SchoolScopedAPITestCase):
    """A family (Amina, her mum, and her dad who also has Brian) with data everywhere."""

    def setUp(self):
        super().setUp()
        self.admin = self.authed_client(self.admin_a)
        year = YearGroup.objects.create(school=self.school_a, name="Grade 7")
        self.klass = SchoolClass.objects.create(year_group=year, name="7A")
        self.assign(self.user_a, self.klass, None)
        self.term = Term.objects.create(school=self.school_a, name="Term 1")
        maths = Subject.objects.create(school=self.school_a, name="Maths")
        self.kid = Student.objects.create(
            school=self.school_a, school_class=self.klass, first_name="Amina", last_name="Otieno",
            external_id="ADM1", date_of_birth=date(2013, 5, 1), medical_notes="Asthma", photo=b"jpeg")
        self.sibling = Student.objects.create(school=self.school_a, school_class=self.klass, first_name="Brian",
                                              last_name="Otieno")
        self.other = Student.objects.create(school=self.school_a, school_class=self.klass, first_name="Cy",
                                            last_name="Kamau")
        for s, score in ((self.kid, 80), (self.other, 40)):
            Grade.objects.create(student=s, subject=maths, term=self.term, score=score)
        AttendanceRecord.objects.create(student=self.kid, date=date(2026, 2, 2), status="excused",
                                        notes="Hospital appointment")
        StudentReport.objects.create(student=self.kid, term=self.term, progress_summary="Staff note",
                                     report_comment="Amina did well", status="finalized")
        # Mum has only Amina; Dad also has Brian at the school.
        self.mum = self.parent("mum@alpha.test", "Grace Otieno", [self.kid], phone="+254 700 111 222",
                               address="12 Ngong Road")
        self.dad = self.parent("dad@alpha.test", "Peter Otieno", [self.kid, self.sibling])
        chat = Conversation.objects.create(school=self.school_a, kind=Conversation.Kind.DIRECT, student=self.kid,
                                           created_by=self.user_a)
        ConversationParticipant.objects.create(conversation=chat, user=self.user_a)
        ConversationParticipant.objects.create(conversation=chat, user=self.mum.user)
        Message.objects.create(conversation=chat, sender=self.mum.user, body="Amina has a cough today")
        Message.objects.create(conversation=chat, sender=self.user_a, body="Thanks for letting me know")
        ActivityLog.objects.create(school=self.school_a, action="grade.updated",
                                   summary="Changed Amina Otieno's Maths grade")
        ActivityLog.objects.create(school=self.school_a, action="parent_invite.accepted",
                                   summary="Parent Grace Otieno accepted their invite")

    def parent(self, email, name, kids, **extra):
        user = User.objects.create_user(username=email, email=email, password="pass1234")
        g = Guardian.objects.create(user=user, school=self.school_a, display_name=name, **extra)
        g.students.set(kids)
        return g

    def remove(self, name="Amina Otieno", client=None):
        return (client or self.admin).post(f"/api/students/{self.kid.id}/remove-personal-data/",
                                           {"confirm_name": name})


class FamilyDataTests(FamilyDataFixture):
    def test_export_contains_everything_held(self):
        response = self.admin.get(f"/api/students/{self.kid.id}/data-export/")
        self.assertEqual(response.status_code, 200)
        self.assertIn("personal-data-amina-otieno.xlsx", response["Content-Disposition"])
        wb = openpyxl.load_workbook(BytesIO(response.content))
        # Other features add sheets after these (applications, roll calls, the change log): they have their own tests.
        self.assertEqual(wb.sheetnames[:10], ["Student", "Parents", "Grades", "Attendance", "Reports", "Messages",
                                              "Subject comments", "About the student", "Support", "Boarding"])
        student = {row[0]: row[1] for row in wb["Student"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual((student["First name"], student["Health notes"], student["Photo held"]),
                         ("Amina", "Asthma", "Yes"))
        parents = list(wb["Parents"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(sorted(p[0] for p in parents), ["Grace Otieno", "Peter Otieno"])
        self.assertEqual(len(list(wb["Messages"].iter_rows(min_row=2))), 2)
        self.assertTrue(ActivityLog.objects.filter(action="student.data_exported").exists())

    def test_only_admins_at_the_school_can_export_or_remove(self):
        self.assertEqual(self.client_a.get(f"/api/students/{self.kid.id}/data-export/").status_code, 403)
        self.assertEqual(self.remove(client=self.client_a).status_code, 403)
        other_admin = User.objects.create_user(username="ob", email="ob@beta.test", password="pass1234")
        Profile.objects.create(user=other_admin, school=self.school_b, role=Profile.Role.ADMIN)
        outsider = self.authed_client(other_admin)
        self.assertEqual(outsider.get(f"/api/students/{self.kid.id}/data-export/").status_code, 404)
        self.assertEqual(self.remove(client=outsider).status_code, 404)

    def test_removal_needs_the_exact_name(self):
        self.assertEqual(self.remove("Amina").status_code, 400)
        self.kid.refresh_from_db()
        self.assertEqual(self.kid.first_name, "Amina")

    def test_removal_clears_personal_details_but_keeps_results(self):
        response = self.remove("  amina   OTIENO ")
        self.assertEqual(response.status_code, 200)
        self.kid.refresh_from_db()
        self.assertEqual((self.kid.first_name, self.kid.last_name), ("Removed", "student"))
        self.assertEqual((self.kid.external_id, self.kid.medical_notes, self.kid.date_of_birth, self.kid.photo),
                         ("", "", None, None))
        self.assertFalse(self.kid.is_active)
        # Results stay so class averages don't change; notes and reports go.
        self.assertEqual(Grade.objects.filter(student=self.kid).count(), 1)
        self.assertEqual(AttendanceRecord.objects.get(student=self.kid).notes, "")
        self.assertFalse(StudentReport.objects.filter(student=self.kid).exists())

    def test_removal_deletes_sole_parents_and_keeps_shared_ones(self):
        self.remove()
        self.assertFalse(User.objects.filter(email="mum@alpha.test").exists())
        self.assertFalse(Message.objects.filter(body__contains="cough").exists())
        self.assertFalse(Conversation.objects.filter(kind=Conversation.Kind.DIRECT).exists())
        self.dad.refresh_from_db()
        self.assertEqual(list(self.dad.students.all()), [self.sibling])
        self.assertTrue(User.objects.get(email="dad@alpha.test").is_active)

    def test_removal_scrubs_names_from_the_activity_log(self):
        self.remove()
        summaries = list(ActivityLog.objects.values_list("summary", flat=True))
        self.assertFalse(any("Amina Otieno" in s or "Grace Otieno" in s for s in summaries))
        self.assertIn("Changed Removed student's Maths grade", summaries)
        self.assertTrue(ActivityLog.objects.filter(action="student.personal_data_removed").exists())

    def test_other_students_are_untouched(self):
        self.remove()
        self.other.refresh_from_db()
        self.sibling.refresh_from_db()
        self.assertEqual((self.other.first_name, self.sibling.first_name), ("Cy", "Brian"))
        self.assertEqual(Grade.objects.filter(student=self.other).count(), 1)
