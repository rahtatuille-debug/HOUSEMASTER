"""The Excel import: preview, import, errors, and school separation."""
from io import BytesIO

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile

from accounts.tests import SchoolScopedAPITestCase
from activity.models import ActivityLog
from attendance.models import AttendanceRecord
from gradebook.models import Grade

from .models import SchoolClass, Student


def workbook(students, grades=None, attendance=None):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in (("Students", students), ("Grades", grades), ("Attendance", attendance)):
        if rows is None:
            continue
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    out = BytesIO()
    wb.save(out)
    return SimpleUploadedFile("school.xlsx", out.getvalue(),
                              content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


STUDENTS = [
    ["id", "first_name", "last_name", "class", "year_group", "gender", "date_of_birth"],
    ["S1", "Ann", "One", "7A", "Year 7", "female", "2014-02-03"],
    ["S2", "Ben", "Two", "7A", "Year 7", "", ""],
    ["S3", "", "Nameless", "7A", "Year 7", "", ""],          # error: no first name
]
GRADES = [
    ["student_id", "subject", "term", "score", "max_score"],
    ["S1", "Maths", "Term 1", 80, 100],
    ["S2", "Maths", "Term 1", 120, 100],                    # error: over max
    ["S9", "Maths", "Term 1", 50, 100],                     # error: unknown student
]
ATTENDANCE = [
    ["student_id", "date", "status", "notes"],
    ["S1", "2026-02-03", "absent", "Ill"],
    ["S2", "2026-02-03", "sleeping", ""],                   # error: bad status
]


class ImportTests(SchoolScopedAPITestCase):
    def setUp(self):
        super().setUp()
        self.admin_client_a = self.authed_client(self.admin_a)

    def post(self, commit, client=None, file=None):
        return (client or self.admin_client_a).post(
            "/api/import/", {"file": file or workbook(STUDENTS, GRADES, ATTENDANCE), "commit": str(commit).lower()},
            format="multipart",
        )

    def test_preview_reports_everything_and_saves_nothing(self):
        response = self.post(commit=False)
        self.assertEqual(response.status_code, 200)
        data = response.data
        self.assertFalse(data["committed"])
        self.assertEqual(data["counts"], {"students_created": 2, "students_updated": 0, "grades": 1, "attendance": 1})
        self.assertEqual(data["created"]["classes"], ["7A (Year 7)"])
        self.assertEqual(data["created"]["subjects"], ["Maths"])
        messages = [(e["sheet"], e["row"]) for e in data["errors"]]
        self.assertEqual(messages, [("Students", 4), ("Grades", 3), ("Grades", 4), ("Attendance", 3)])
        self.assertFalse(Student.objects.filter(school=self.school_a).exists())
        self.assertFalse(SchoolClass.objects.filter(name="7A").exists())

    def test_import_saves_the_good_rows_and_logs_it(self):
        response = self.post(commit=True)
        self.assertTrue(response.data["committed"])
        ann = Student.objects.get(school=self.school_a, external_id="S1")
        self.assertEqual((ann.school_class.name, ann.school_class.year_group.name, ann.gender), ("7A", "Year 7", "female"))
        self.assertEqual(Grade.objects.get(student=ann).score, 80)
        self.assertEqual(AttendanceRecord.objects.get(student=ann).status, "absent")
        self.assertFalse(Grade.objects.filter(student__external_id="S2").exists())
        self.assertTrue(ActivityLog.objects.filter(action="school.imported").exists())

    def test_reimporting_updates_instead_of_duplicating(self):
        self.post(commit=True)
        response = self.post(commit=True)
        self.assertEqual(response.data["counts"]["students_created"], 0)
        self.assertEqual(response.data["counts"]["students_updated"], 2)
        self.assertEqual(Student.objects.filter(school=self.school_a).count(), 2)

    def test_same_admission_number_at_another_school_is_never_touched(self):
        other = Student.objects.create(school=self.school_b, external_id="S1", first_name="Other", last_name="Kid")
        self.post(commit=True)
        other.refresh_from_db()
        self.assertEqual((other.first_name, other.school), ("Other", self.school_b))
        self.assertEqual(Student.objects.filter(external_id="S1").count(), 2)

    def test_existing_class_is_reused(self):
        from .models import YearGroup

        year = YearGroup.objects.create(school=self.school_a, name="Year 7")
        existing = SchoolClass.objects.create(year_group=year, name="7A")
        data = self.post(commit=True).data
        self.assertEqual(data["created"]["classes"], [])
        self.assertEqual(Student.objects.get(external_id="S1").school_class, existing)

    def test_teacher_cannot_import(self):
        self.assertEqual(self.post(commit=True, client=self.client_a).status_code, 403)
        self.assertEqual(self.client_a.get("/api/import/template/").status_code, 403)

    def test_not_an_excel_file(self):
        bad = SimpleUploadedFile("x.xlsx", b"hello", content_type="application/octet-stream")
        response = self.post(commit=False, file=bad)
        self.assertEqual(response.status_code, 400)

    def test_missing_students_sheet(self):
        wb = openpyxl.Workbook()
        wb.active.title = "Sheet1"
        out = BytesIO()
        wb.save(out)
        response = self.post(commit=False, file=SimpleUploadedFile("x.xlsx", out.getvalue()))
        self.assertEqual(response.status_code, 400)
        self.assertIn("Students", str(response.data))

    def test_template_download_is_a_valid_workbook(self):
        response = self.admin_client_a.get("/api/import/template/")
        self.assertEqual(response.status_code, 200)
        wb = openpyxl.load_workbook(BytesIO(response.content))
        self.assertEqual(wb.sheetnames, ["Students", "Grades", "Attendance"])
        # The template's own example rows import cleanly.
        result = self.post(commit=False, file=SimpleUploadedFile("t.xlsx", response.content)).data
        self.assertEqual(result["errors"], [])
