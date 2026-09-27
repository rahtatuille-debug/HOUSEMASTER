"""Setup's people steps: staff, students and parents must be added before the school opens."""
from io import BytesIO

import openpyxl

from accounts.models import Invite
from accounts.tests import SchoolScopedAPITestCase
from guardians.models import ClassSignupLink

from .models import Student
from .test_setup import finish_payload


class PeopleStepsTests(SchoolScopedAPITestCase):
    def setUp(self):
        super().setUp()
        self.admin = self.authed_client(self.admin_a)
        self.admin.post("/api/setup/finish/", finish_payload(), format="json")

    def status(self):
        return self.admin.get("/api/setup/people/").data

    def test_the_school_opens_only_once_staff_students_and_parents_are_in(self):
        # teacher_a already works here, so staff is done; the other two aren't.
        self.assertEqual(self.status()["stage"], "people")
        response = self.admin.post("/api/setup/complete/")
        self.assertEqual(response.status_code, 400)
        self.assertIn("add your students", str(response.data))
        self.assertIn("sign-up link", str(response.data))

        from .models import SchoolClass
        klass = SchoolClass.objects.filter(year_group__school=self.school_a).first()
        Student.objects.create(school=self.school_a, school_class=klass, first_name="Amina", last_name="O")
        links = self.admin.post("/api/signup-links/", {"action": "create_all"}, format="json").data
        self.assertTrue(all(row["token"] for row in links))
        self.assertEqual(ClassSignupLink.objects.filter(school=self.school_a).count(), 3)

        status = self.status()
        self.assertTrue(status["staff"]["done"] and status["students"]["done"] and status["parents"]["done"])
        self.assertEqual(self.admin.post("/api/setup/complete/").data["stage"], "done")
        school = self.admin.get("/api/me/").data["school"]
        self.assertEqual((school["setup_completed"], school["setup_stage"]), (True, "done"))

    def test_staff_means_someone_besides_the_admin(self):
        self.user_a.profile.delete()
        self.assertFalse(self.status()["staff"]["done"])
        Invite.objects.create(school=self.school_a, email="new@alpha.test", name="New Teacher", role="teacher",
                              invited_by=self.admin_a)
        self.assertTrue(self.status()["staff"]["done"])

    def test_the_structure_comes_first(self):
        self.school_b.refresh_from_db()
        self.make_admin(self.user_b)
        self.assertEqual(self.client_b.get("/api/setup/people/").data["stage"], "structure")
        self.assertEqual(self.client_b.post("/api/setup/complete/").status_code, 400)
        self.assertEqual(self.client_a.post("/api/setup/complete/").status_code, 403)

    def test_templates_use_the_schools_own_classes_and_subjects(self):
        students = openpyxl.load_workbook(BytesIO(self.admin.get("/api/import/template/").content))
        self.assertEqual(students["Students"]["D2"].value, "7 East")
        self.assertEqual([r[0] for r in students["Classes"].iter_rows(min_row=2, values_only=True)],
                         ["7 East", "7 West", "Grade 8"])
        staff = openpyxl.load_workbook(BytesIO(self.admin.get("/api/import/staff-template/").content))
        self.assertEqual(staff["Staff"]["D2"].value, "7 East")
        self.assertIn("English", staff["Staff"]["E2"].value)
        self.assertEqual({r[1] for r in staff["Classes and subjects"].iter_rows(min_row=2, values_only=True)} - {None},
                         {"English", "Mathematics"})
