"""
Import a school's Excel workbook (Students / Grades / Attendance sheets).

Used by the in-app import (admins) and the import_school_workbook command.
Preview and import run exactly the same code: a preview does the whole
import inside a transaction and then rolls it back, so what the preview
reports is what the import will do.

Sheets and columns (first row is the header, column order doesn't matter):
  Students:   id, first_name, last_name, class, house, year_group,
              gender, date_of_birth
  Grades:     student_id, subject, term, score, max_score
  Attendance: student_id, date, status, notes

`id` / `student_id` is the school's own admission number (Student.external_id).
Only the Students sheet is required.
"""
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

import openpyxl
from django.db import transaction

from attendance.models import AttendanceRecord
from gradebook.models import Grade, Subject, Term

from .models import SchoolClass, Student, YearGroup

MAX_ERRORS = 200
TEMPLATE = {
    "Students": ["id", "first_name", "last_name", "class", "year_group", "house", "gender", "date_of_birth"],
    "Grades": ["student_id", "subject", "term", "score", "max_score"],
    "Attendance": ["student_id", "date", "status", "notes"],
}
TEMPLATE_EXAMPLES = {
    "Students": ["BS2068", "John", "Doe", "7A", "Year 7", "Kilimanjaro", "male", "2014-01-31"],
    "Grades": ["BS2068", "Mathematics", "Term 1 2026", 78, 100],
    "Attendance": ["BS2068", "2026-02-03", "present", ""],
}
STATUSES = {c[0] for c in AttendanceRecord.STATUS_CHOICES}
GENDERS = {c[0] for c in Student.Gender.choices}


class WorkbookError(Exception):
    """The file can't be imported at all (not an Excel file, no Students sheet...)."""


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # Excel stores 2068 as 2068.0
    return str(value).strip()


def _date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ValueError(f'"{text}" isn\'t a date (use YYYY-MM-DD)')


def _rows(ws):
    """Yield (row number, dict keyed by lower-case header) for each non-empty row."""
    headers = [_text(c.value).lower() for c in ws[1]]
    for number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(v is None or _text(v) == "" for v in row):
            continue
        yield number, dict(zip(headers, row))


class _Import:
    def __init__(self, school):
        self.school = school
        self.errors = []
        self.counts = {"students_created": 0, "students_updated": 0, "grades": 0, "attendance": 0}
        self.created = {"year_groups": [], "classes": [], "subjects": [], "terms": []}

    def error(self, sheet, row, message):
        if len(self.errors) < MAX_ERRORS:
            self.errors.append({"sheet": sheet, "row": row, "message": message})

    def _year_group(self, name):
        year_group, created = YearGroup.objects.get_or_create(school=self.school, name=name)
        if created:
            self.created["year_groups"].append(name)
        return year_group

    def _class(self, name, year_group_name):
        existing = SchoolClass.objects.filter(year_group__school=self.school, name=name)
        if year_group_name:
            existing = existing.filter(year_group__name=year_group_name)
        match = existing.first()
        if match:
            return match
        year_group = self._year_group(year_group_name or "Imported")
        school_class = SchoolClass.objects.create(year_group=year_group, name=name)
        self.created["classes"].append(f"{name} ({year_group.name})")
        return school_class

    def _named(self, model, name, bucket):
        obj, created = model.objects.get_or_create(school=self.school, name=name)
        if created:
            self.created[bucket].append(name)
        return obj

    def _student(self, sheet, number, ext_id):
        if not ext_id:
            self.error(sheet, number, "student_id is empty")
            return None
        student = Student.objects.filter(school=self.school, external_id=ext_id).first()
        if student is None:
            self.error(sheet, number, f'No student with admission number "{ext_id}"')
        return student

    def students(self, ws):
        for number, row in _rows(ws):
            ext_id, first, last = _text(row.get("id")), _text(row.get("first_name")), _text(row.get("last_name"))
            if not ext_id or not first or not last:
                self.error("Students", number, "id, first_name and last_name are all required")
                continue
            gender = _text(row.get("gender")).lower()
            if gender and gender not in GENDERS:
                self.error("Students", number, f'gender "{gender}" must be female, male or other')
                continue
            try:
                dob = _date(row.get("date_of_birth"))
            except ValueError as exc:
                self.error("Students", number, f"date_of_birth: {exc}")
                continue
            class_name = _text(row.get("class"))
            values = {"first_name": first, "last_name": last, "house": _text(row.get("house"))}
            if class_name:
                values["school_class"] = self._class(class_name, _text(row.get("year_group")))
            if gender:
                values["gender"] = gender
            if dob:
                values["date_of_birth"] = dob
            _, created = Student.objects.update_or_create(school=self.school, external_id=ext_id, defaults=values)
            self.counts["students_created" if created else "students_updated"] += 1

    def grades(self, ws):
        for number, row in _rows(ws):
            student = self._student("Grades", number, _text(row.get("student_id")))
            subject_name, term_name = _text(row.get("subject")), _text(row.get("term"))
            if not subject_name or not term_name:
                self.error("Grades", number, "subject and term are required")
                continue
            try:
                score = Decimal(_text(row.get("score")))
                max_raw = _text(row.get("max_score"))
                max_score = Decimal(max_raw) if max_raw else Decimal("100")
            except InvalidOperation:
                self.error("Grades", number, "score and max_score must be numbers")
                continue
            if max_score <= 0 or score < 0 or score > max_score:
                self.error("Grades", number, f"score {score} must be between 0 and {max_score}")
                continue
            if student is None:
                continue
            term = self._named(Term, term_name, "terms")
            if getattr(term, "is_locked", False):
                self.error("Grades", number, f'Term "{term_name}" is locked')
                continue
            subject = self._named(Subject, subject_name, "subjects")
            Grade.objects.update_or_create(student=student, subject=subject, term=term,
                                           defaults={"score": score, "max_score": max_score})
            self.counts["grades"] += 1

    def attendance(self, ws):
        for number, row in _rows(ws):
            student = self._student("Attendance", number, _text(row.get("student_id")))
            try:
                day = _date(row.get("date"))
            except ValueError as exc:
                self.error("Attendance", number, f"date: {exc}")
                continue
            if day is None:
                self.error("Attendance", number, "date is required")
                continue
            status = _text(row.get("status")).lower() or "present"
            if status not in STATUSES:
                self.error("Attendance", number, f'status "{status}" must be present, absent, late or excused')
                continue
            if student is None:
                continue
            from gradebook.locks import locked_term_for_date

            locked = locked_term_for_date(self.school, day)
            if locked is not None:
                self.error("Attendance", number, f"{day} is in {locked.name}, which is locked")
                continue
            AttendanceRecord.objects.update_or_create(
                student=student, date=day, defaults={"status": status, "notes": _text(row.get("notes"))[:255]}
            )
            self.counts["attendance"] += 1


def import_workbook(file, school, commit):
    """
    Import (commit=True) or preview (commit=False) a workbook for one school.
    Returns a summary: counts, what was/would be created, and row-by-row errors.
    Rows with errors are skipped; the rest are imported.
    """
    try:
        wb = openpyxl.load_workbook(file, data_only=True, read_only=False)
    except Exception:
        raise WorkbookError("That file isn't an Excel workbook (.xlsx) HouseMaster can read.")
    if "Students" not in wb.sheetnames:
        raise WorkbookError('The workbook needs a sheet called "Students". Download the template to see the layout.')

    job = _Import(school)
    with transaction.atomic():
        job.students(wb["Students"])
        if "Grades" in wb.sheetnames:
            job.grades(wb["Grades"])
        if "Attendance" in wb.sheetnames:
            job.attendance(wb["Attendance"])
        if not commit:
            transaction.set_rollback(True)
    return {
        "committed": commit,
        "counts": job.counts,
        "created": job.created,
        "errors": job.errors,
        "errors_truncated": len(job.errors) >= MAX_ERRORS,
    }


def template_workbook(school=None):
    """
    An empty workbook with the right sheets, headers and one example row each.
    For a school, the example uses its own classes, subjects and terms, and a
    "Classes" sheet lists every class name to copy from.
    """
    examples = {k: list(v) for k, v in TEMPLATE_EXAMPLES.items()}
    classes = []
    if school is not None:
        classes = list(SchoolClass.objects.filter(year_group__school=school).select_related("year_group")
                       .order_by("year_group__order", "year_group__name", "name"))
        if classes:
            examples["Students"][3], examples["Students"][4] = classes[0].name, classes[0].year_group.name
        subject = Subject.objects.filter(school=school).order_by("name").first()
        term = Term.objects.filter(school=school).order_by("start_date", "name").first()
        if subject:
            examples["Grades"][1] = subject.name
        if term:
            examples["Grades"][2] = term.name
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet, headers in TEMPLATE.items():
        ws = wb.create_sheet(sheet)
        ws.append(headers)
        ws.append(examples[sheet])
        for cell in ws[1]:
            cell.font = openpyxl.styles.Font(bold=True)
        for column, header in zip(ws.columns, headers):
            ws.column_dimensions[column[0].column_letter].width = max(12, len(header) + 4)
    if classes:
        ws = wb.create_sheet("Classes")
        ws.append(["class", "year_group"])
        for klass in classes:
            ws.append([klass.name, klass.year_group.name])
        for cell in ws[1]:
            cell.font = openpyxl.styles.Font(bold=True)
        ws.column_dimensions["A"].width = ws.column_dimensions["B"].width = 22
    return wb
